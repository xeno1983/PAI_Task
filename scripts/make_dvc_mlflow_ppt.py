"""Build a short PPT explaining Git / MLflow / DVC / MinIO with real UI screenshots."""

from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

FONT = "Yu Gothic"
INK = RGBColor(0x1C, 0x28, 0x2C)
MUTED = RGBColor(0x4A, 0x58, 0x5C)
PAPER = RGBColor(0xF6, 0xF3, 0xEC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEAL = RGBColor(0x1F, 0x4E, 0x5F)
GREEN = RGBColor(0x2C, 0x64, 0x4A)
CLAY = RGBColor(0x8A, 0x45, 0x2A)
SOFT = RGBColor(0xE8, 0xE4, 0xDC)

ASSETS = Path(__file__).resolve().parent / "assets" / "dvc_mlflow_ppt"
LAUNCHER = Path.home() / "Desktop" / "起動スクリプト" / "PickAndCarry"


def paint(run, size, bold=False, color=INK) -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    props = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        node = props.find(qn(tag))
        if node is None:
            node = props.makeelement(qn(tag), {})
            props.append(node)
        node.set("typeface", FONT)


def write(shape, lines, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT, spacing=6) -> None:
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    for index, line in enumerate(lines):
        paragraph = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        paragraph.alignment = align
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(spacing)
        run = paragraph.add_run()
        run.text = line
        paint(run, size, bold=bold, color=color)


def blank(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = PAPER
    bg.line.fill.background()
    return slide


def title_bar(slide, title: str, subtitle: str | None = None) -> None:
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.95))
    bar.fill.solid()
    bar.fill.fore_color.rgb = TEAL
    bar.line.fill.background()
    box = slide.shapes.add_textbox(Inches(0.45), Inches(0.18), Inches(12.4), Inches(0.7))
    write(box, [title], size=26, bold=True, color=WHITE, spacing=0)
    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.45), Inches(1.05), Inches(12.4), Inches(0.4))
        write(sub, [subtitle], size=13, color=MUTED, spacing=0)


def add_picture_fit(slide, path: Path, left, top, width, height) -> None:
    if not path.is_file():
        missing = slide.shapes.add_textbox(left, top, width, height)
        write(missing, [f"(missing screenshot: {path.name})"], size=12, color=CLAY)
        return
    frame = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    frame.adjustments[0] = 0.04
    frame.fill.solid()
    frame.fill.fore_color.rgb = WHITE
    frame.line.color.rgb = SOFT
    pad = Inches(0.1)
    box_w = float(width - 2 * pad)
    box_h = float(height - 2 * pad)
    with Image.open(path) as img:
        iw, ih = img.size
    scale = min(box_w / iw, box_h / ih)
    pic_w = iw * scale
    pic_h = ih * scale
    pic_left = left + pad + (box_w - pic_w) / 2
    pic_top = top + pad + (box_h - pic_h) / 2
    slide.shapes.add_picture(str(path), pic_left, pic_top, width=pic_w, height=pic_h)


def build() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1 title
    s = blank(prs)
    hero = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    hero.fill.solid()
    hero.fill.fore_color.rgb = TEAL
    hero.line.fill.background()
    box = s.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.5), Inches(2.5))
    write(
        box,
        [
            "Git × MLflow × DVC × MinIO",
            "Pick-and-Carry での実験記録と重み共有",
        ],
        size=32,
        bold=True,
        color=WHITE,
        spacing=12,
    )
    foot = s.shapes.add_textbox(Inches(0.8), Inches(5.4), Inches(11.5), Inches(1.0))
    write(
        foot,
        [
            "実機の画面（MinIO Console / MLflow UI / DVC status）を使った使い方まとめ",
            "ログイン: MinIO Console http://127.0.0.1:9001  /  MLflow http://127.0.0.1:5000",
        ],
        size=14,
        color=RGBColor(0xD7, 0xE4, 0xE8),
        spacing=4,
    )

    # 2 architecture
    s = blank(prs)
    title_bar(s, "全体像：役割分担", "学習を速くする道具ではなく、あとから困らないための棚分け")
    roles = [
        (Inches(0.45), "Git", GREEN, ["コード・設定", ".dvc ポインタ", "設計図"]),
        (Inches(3.55), "MLflow", TEAL, ["パラメータ", "メトリクス曲線", "実験ノート"]),
        (Inches(6.65), "DVC", CLAY, ["大きい重みの版", "best.pt の選定", "荷物の伝票"]),
        (Inches(9.75), "MinIO", MUTED, ["S3 互換倉庫", "pai-dvc / pai-mlflow", "実ファイル"]),
    ]
    for x, name, accent, lines in roles:
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.7), Inches(2.9), Inches(3.6))
        card.adjustments[0] = 0.08
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.fill.background()
        head = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, Inches(1.7), Inches(2.9), Inches(0.7))
        head.fill.solid()
        head.fill.fore_color.rgb = accent
        head.line.fill.background()
        write(head, [name], size=20, bold=True, color=WHITE, align=PP_ALIGN.CENTER, spacing=0)
        body = s.shapes.add_textbox(x + Inches(0.2), Inches(2.6), Inches(2.5), Inches(2.4))
        write(body, lines, size=16, color=INK, spacing=10)
    note = s.shapes.add_textbox(Inches(0.45), Inches(5.6), Inches(12.4), Inches(1.2))
    write(
        note,
        [
            "学習中の生データは logs/（作業机）。区切りだけ DVC publish → MinIO push。",
            "MinIO は DB ではない。オブジェクトストレージ（ファイル倉庫）です。",
        ],
        size=14,
        color=MUTED,
        spacing=4,
    )

    # 3 flow
    s = blank(prs)
    title_bar(s, "いつ何を使うか", "ホットパスでは使わず、区切り・引っ越し・提出で使う")
    steps = [
        "1. 学習 → logs/<stage>/ppo_*.pt + metrics.jsonl + MLflow 記録",
        "2. 区切り → メニュー 9 → 3（最良）または 4（最新）で publish → push",
        "3. 確認 → MinIO Console の pai-dvc（実体） / Git の .dvc（ポインタ）",
        "4. 復元・推論 → pull → restore、または 9 → 6 で best.pt を play",
        "5. 曲線確認 → MLflow UI（実験ノート）。重み本体は DVC/MinIO",
    ]
    box = s.shapes.add_textbox(Inches(0.7), Inches(1.8), Inches(12), Inches(4.5))
    write(box, steps, size=20, color=INK, spacing=16)

    # 4 MinIO login
    s = blank(prs)
    title_bar(s, "MinIO Console：ログイン", "http://127.0.0.1:9001   user / pass = minioadmin（.env と同じ）")
    add_picture_fit(s, ASSETS / "minio_login.png", Inches(1.2), Inches(1.6), Inches(10.8), Inches(5.4))

    # 5 MinIO buckets
    s = blank(prs)
    title_bar(s, "MinIO：バケット一覧（実画面）", "pai-dvc = DVC の重み倉庫　／　pai-mlflow = MLflow artifact 用")
    add_picture_fit(s, ASSETS / "minio_pai_dvc.png", Inches(0.5), Inches(1.55), Inches(12.3), Inches(5.5))

    # 6 MinIO hashed files
    s = blank(prs)
    title_bar(s, "MinIO の中身：DVC の content-addressed 保管", "pai-dvc / files / md5/ … ハッシュ名で実体を置く（Git には載せない）")
    add_picture_fit(s, ASSETS / "minio_pai_dvc_files.png", Inches(0.5), Inches(1.55), Inches(12.3), Inches(5.5))

    # 7 DVC
    s = blank(prs)
    title_bar(s, "DVC：ポインタと最良選定", "Git には .dvc、実体は MinIO。publish --select best で metrics から選ぶ")
    add_picture_fit(s, ASSETS / "dvc_status.png", Inches(0.5), Inches(1.55), Inches(7.4), Inches(5.2))
    side = s.shapes.add_textbox(Inches(8.2), Inches(1.8), Inches(4.6), Inches(4.8))
    write(
        side,
        [
            "メニュー 9",
            "3  最良を publish → push",
            "4  最新を publish → push",
            "6  最良で推論（best.pt）",
            "",
            "CLI 例",
            "dvc_sync.py publish \\",
            "  --log-dir logs\\... \\",
            "  --select best --push",
            "",
            "resolve で best.pt のパス表示",
        ],
        size=15,
        color=INK,
        spacing=6,
    )

    # 8 MLflow home
    s = blank(prs)
    title_bar(s, "MLflow：実験ノート（実画面）", "起動: mlflow ui --backend-store-uri sqlite:///logs/mlflow/mlflow.db")
    add_picture_fit(s, ASSETS / "mlflow_home.png", Inches(0.5), Inches(1.55), Inches(12.3), Inches(5.5))

    # 9 MLflow runs
    s = blank(prs)
    title_bar(s, "MLflow：pick-and-carry の Runs", "曲線・パラメータ比較はここ。大きい .pt の共有は DVC/MinIO 側")
    add_picture_fit(s, ASSETS / "mlflow_experiment.png", Inches(0.5), Inches(1.55), Inches(12.3), Inches(5.5))

    # 10 checklist
    s = blank(prs)
    title_bar(s, "最短チェックリスト", "起動スクリプト → メニュー 9")
    lines = [
        "□ 9 → 1  MinIO 起動（Docker Desktop も必要なら起動）",
        "□ 9 → 2  初回のみ setup（pip + DVC remote）",
        "□ 9 → 3  学習済み段階の最良重みを publish → push",
        "□ Console で pai-dvc に objects が増えたことを確認",
        "□ 9 → 6  または resolve + --play で最良を推論",
        "□ MLflow UI で曲線を見る（メトリクスは別系統）",
        "",
        "詳細手順: docs/DVC_MINIO.md",
    ]
    box = s.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.5), Inches(5.0))
    write(box, lines, size=20, color=INK, spacing=12)

    return prs


def main() -> None:
    ASSETS.mkdir(parents=True, exist_ok=True)
    prs = build()
    out_repo = Path(__file__).resolve().parent / "DVC_MLflow_MinIO使い方.pptx"
    out_desk = LAUNCHER / "DVC_MLflow_MinIO使い方.pptx"
    LAUNCHER.mkdir(parents=True, exist_ok=True)
    prs.save(out_repo)
    prs.save(out_desk)
    print(f"saved {out_repo}")
    print(f"saved {out_desk}")


if __name__ == "__main__":
    main()
